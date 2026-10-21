"""Groq-powered, read-only clinic information agents.

Appointment creation and cancellation remain in chatbot.py, where the clinic's
existing validation, confirmation, and slot-conflict handling are enforced.
"""
import json
import os
import re
from functools import lru_cache
from pathlib import Path

from database import all_doctors, free_slots
from chatbot import CLINIC


AGENTS = {
    "doctors": {
        "name": "Doctors and fees agent",
        "instructions": "Help visitors compare clinic doctors, specialties, and listed fees. Use the doctor tool for factual answers.",
    },
    "clinic_info": {
        "name": "Clinic information agent",
        "instructions": "Answer questions about clinic hours, address, phone, and preparing for a visit. Use the clinic information tool for facts.",
    },
}

TOOLS = [
    {"type": "function", "function": {
        "name": "list_doctors",
        "description": "Get active clinic doctors, specialties, and consultation fees.",
        "parameters": {"type": "object", "properties": {}, "required": [], "additionalProperties": False}}},
    {"type": "function", "function": {
        "name": "get_clinic_information",
        "description": "Get the clinic's verified address, phone number, and opening hours.",
        "parameters": {"type": "object", "properties": {}, "required": [], "additionalProperties": False}}},
    {"type": "function", "function": {
        "name": "get_available_slots",
        "description": "Check sample appointment slots for a doctor on an ISO date (YYYY-MM-DD).",
        "parameters": {"type": "object", "properties": {
            "doctor_id": {"type": "integer"}, "date": {"type": "string"}},
            "required": ["doctor_id", "date"], "additionalProperties": False}}},
]


def _execute_tool(name, arguments):
    if name == "list_doctors":
        return [{"id": d["id"], "name": d["name"], "specialty": d["spec"], "fee_inr": d["fee"]}
                for d in all_doctors()]
    if name == "get_clinic_information":
        return {"name": CLINIC["name"], "address": CLINIC["address"],
                "phone": CLINIC["phone"], "hours": CLINIC["hours"], "sunday": "Closed"}
    if name == "get_available_slots":
        day = arguments["date"]
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", day):
            return {"error": "Use a date in YYYY-MM-DD format."}
        try:
            return {"doctor_id": arguments["doctor_id"], "date": day,
                    "sample_slots": free_slots(arguments["doctor_id"], day)}
        except Exception:
            return {"error": "No active doctor matches that ID."}
    return {"error": "Tool is not available."}


@lru_cache(maxsize=1)
def _read_clinic_guide(path, modified_time):
    """Extract the clinic guide's text and tables for the agents' context."""
    try:
        from docx import Document
        document = Document(path)
    except Exception as exc:
        print(f"Clinic guide could not be loaded: {type(exc).__name__}")
        return ""
    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip().replace("\n", "; ") for cell in row.cells]
            parts.append(" | ".join(value for value in values if value))
    return "\n".join(parts)


def _clinic_guide_context():
    guide_path = Path(__file__).with_name("CityCare_Clinic_Information_Guide.docx")
    try:
        modified_time = guide_path.stat().st_mtime_ns
    except OSError:
        return ""
    return _read_clinic_guide(str(guide_path), modified_time)


def _langchain_tools():
    """Build the clinic's read-only tools for the LangChain agents."""
    from langchain_core.tools import tool

    @tool
    def list_doctors() -> str:
        """Get active clinic doctors, their listed specialties, and consultation fees."""
        return json.dumps(_execute_tool("list_doctors", {}), ensure_ascii=False)

    @tool
    def get_clinic_information() -> str:
        """Get the clinic's verified address, telephone number, and opening hours."""
        return json.dumps(_execute_tool("get_clinic_information", {}), ensure_ascii=False)

    @tool
    def get_available_slots(doctor_id: int, date: str) -> str:
        """Check sample appointment slots for a doctor on a date in YYYY-MM-DD format."""
        return json.dumps(_execute_tool("get_available_slots", {
            "doctor_id": doctor_id, "date": date}), ensure_ascii=False)

    return [list_doctors, get_clinic_information, get_available_slots]


def _key_from_secrets():
    try:
        import streamlit as st
        return st.secrets.get("GROQ_API_KEY") or st.secrets.get("groq", {}).get("api_key")
    except Exception:
        return None


def _dotenv_value(setting):
    """Read a single setting from this app's ignored .env file."""
    env_file = os.path.join(os.path.dirname(__file__), ".env")
    try:
        with open(env_file, encoding="utf-8") as source:
            for line in source:
                key, separator, value = line.partition("=")
                if separator and key.strip() == setting:
                    return value.strip().strip('"').strip("'")
    except OSError:
        pass
    return None


def agent_reply(agent_id, history):
    """Run the selected information agent with clinic data tools."""
    agent = AGENTS.get(agent_id)
    if not agent:
        return {"reply": "Choose a clinic assistant option to continue.", "actions": [], "provider": "local"}
    api_key = os.getenv("GROQ_API_KEY") or _dotenv_value("GROQ_API_KEY") or _key_from_secrets()
    if not api_key:
        return {"reply": "The AI agent needs a Groq API key. Add `GROQ_API_KEY=your-key` to `app/.env`, then restart the app. Keep the key private.", "actions": [], "provider": "setup"}

    # Avoid sending common booking identifiers or phone numbers to the model.
    safe_history = []
    for item in history[-8:]:
        content = re.sub(r"\bCC[\s-]*\d{4,6}\b", "[booking reference]", item["text"], flags=re.I)
        content = re.sub(r"(?<!\d)(?:\+?91[\s-]?)?[6-9]\d{9}(?!\d)", "[phone number]", content)
        safe_history.append({"role": item["role"], "content": content})

    try:
        from langchain.agents import create_agent
        from langchain_groq import ChatGroq
        model = os.getenv("GROQ_MODEL") or _dotenv_value("GROQ_MODEL") or "openai/gpt-oss-20b"
        guide = _clinic_guide_context()
        instructions = (f"You are the {agent['name']} for CityCare Clinic in Bhopal. "
                        f"{agent['instructions']} Answer warmly and concisely. Use tools for clinic facts; "
                        "never invent doctors, fees, hours, or availability. This is not a medical advice service. "
                        "For symptoms, diagnosis, or treatment questions, say you cannot advise and recommend contacting a qualified clinician. "
                        "Appointment slots are sample prototype data. Never create, change, or cancel a booking. "
                        "Use the following clinic information guide for visitor-facing clinic details. If its information conflicts "
                        "with a live read-only tool result, use the tool result for current doctor or availability data. "
                        "Do not mention or expose the section titled 'Information to confirm before public use'; it is an internal checklist, not clinic policy.\n\n"
                        f"CLINIC INFORMATION GUIDE\n{guide if guide else 'Guide unavailable; use only the clinic tools and say when details are not listed.'}")
        chat_model = ChatGroq(model=model, groq_api_key=api_key, temperature=0, max_retries=2)
        agent = create_agent(model=chat_model, tools=_langchain_tools(), system_prompt=instructions)
        result = agent.invoke({"messages": safe_history}, config={"recursion_limit": 8})
        content = result["messages"][-1].content
        if isinstance(content, list):
            answer = "\n".join(part.get("text", "") for part in content if isinstance(part, dict)).strip()
        else:
            answer = str(content or "").strip()
        answer = answer or "I couldn't prepare an answer just now. Please call the clinic for help."
        return {"reply": answer, "actions": [], "provider": "groq"}
    except Exception as exc:
        # Don't expose SDK internals or credentials in the visitor-facing UI.
        diagnostic = str(exc).replace(api_key, "[REDACTED]") if api_key else str(exc)
        print(f"Clinic AI agent error ({type(exc).__name__}): {diagnostic}")
        return {"reply": "The AI assistant could not connect to Groq. Check GROQ_API_KEY and GROQ_MODEL in app/.env, then try again.",
                "actions": [], "provider": "error"}
